import json,tempfile,unittest
from pathlib import Path
from openpyxl import Workbook
from demo import make_sample,read_export,publish_local,run,digest,ValidationError

class DemoTests(unittest.TestCase):
    def setUp(self):
        self.temp_base=Path(__file__).resolve().parent/'.test_tmp'
        self.temp_base.mkdir(exist_ok=True)
        self.tmp=tempfile.TemporaryDirectory(dir=self.temp_base);self.root=Path(self.tmp.name)
    def tearDown(self):
        assert self.temp_base.resolve() in self.root.resolve().parents
        self.tmp.cleanup()
    def test_cleaning_preserves_blank_ids_and_repeated_names(self):
        p=self.root/'sample.xlsx';make_sample(p);rows=read_export(p)
        self.assertEqual(len(rows),3);self.assertEqual(rows[2]['ID'],'')
        self.assertEqual(rows[0]['Material'],rows[2]['Material'])
    def test_duplicate_blocks_existing_destination(self):
        run(self.root,apply=True);target=self.root/'google_sheets_simulado.json';before=digest(target)
        with self.assertRaises(ValidationError):run(self.root,'duplicado',True)
        self.assertEqual(digest(target),before)
    def test_conference_does_not_publish(self):
        self.assertFalse(run(self.root)['destino_alterado'])
        self.assertFalse((self.root/'google_sheets_simulado.json').exists())
    def test_backup_matches_previous_content(self):
        run(self.root,apply=True);target=self.root/'google_sheets_simulado.json';before=digest(target)
        run(self.root,apply=True);backups=list((self.root/'backups').glob('*.json'))
        self.assertEqual(len(backups),1);self.assertEqual(digest(backups[0]),before)
    def test_external_change_blocks_write(self):
        target=self.root/'target.json';target.write_text('{}',encoding='utf-8');before=digest(target)
        target.write_text('{"alterado":true}',encoding='utf-8')
        with self.assertRaises(ValidationError):publish_local([],target,before)
        self.assertEqual(json.loads(target.read_text()),{'alterado':True})
    def test_unexpected_headers_and_formula_rejected(self):
        p=self.root/'bad.xlsx';b=Workbook();b.active.append(['Outro cabeçalho']);b.save(p)
        with self.assertRaises(ValidationError):read_export(p)
        make_sample(p);from openpyxl import load_workbook
        b=load_workbook(p);b.active['B3']='=1+1';b.save(p);b.close()
        with self.assertRaises(ValidationError):read_export(p)
    def test_collector_refuses_external_destinations(self):
        from collect_demo import collect
        for url in ['https://monday.com','http://example.com','file:///tmp/board.html']:
            with self.assertRaises(ValueError):collect(url,self.root/'export.xlsx')

if __name__=='__main__':unittest.main()
