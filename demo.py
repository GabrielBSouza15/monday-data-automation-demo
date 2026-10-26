"""Demonstração independente: XLSX fictício -> validação -> destino JSON local.

Não lê configurações, credenciais, perfis ou bases de nenhum projeto operacional.
O JSON representa o destino tabular; não é uma integração real com Google Sheets.
"""
from __future__ import annotations
import argparse
from datetime import datetime, timezone
import hashlib
import json
from pathlib import Path
import tempfile
from openpyxl import Workbook, load_workbook

COLUMNS=('ID','Material','Etapa','Responsável','Prazo')
SAMPLE=[['DEMO-001','Guia de astronomia','Revisão','Pessoa fictícia A','2026-11-02'],['DEMO-002','Caderno de ecologia','Diagramação','Pessoa fictícia B','2026-11-05'],['','Guia de astronomia','Planejamento','Pessoa fictícia C','2026-11-08']]

class ValidationError(ValueError):pass

def digest(path:Path)->str:
    return hashlib.sha256(path.read_bytes()).hexdigest()

def make_sample(path:Path, duplicate=False):
    book=Workbook(); sheet=book.active; sheet.title='Quadro fictício'
    sheet.append(['Grupo: materiais de demonstração']);sheet.append(list(COLUMNS))
    for row in SAMPLE:sheet.append(row)
    sheet.append([None]*5);sheet.append(list(COLUMNS))
    if duplicate:sheet.append(['DEMO-001','Outro registro fictício','Revisão','Pessoa fictícia D','2026-11-10'])
    sheet.append(['Resumo do grupo',None,None,None,None])
    path.parent.mkdir(parents=True,exist_ok=True);book.save(path)

def read_export(path:Path)->list[dict]:
    book=load_workbook(path,read_only=True,data_only=False)
    records=[]; header_seen=False
    try:
        for row in book.active.iter_rows(values_only=True):
            values=list(row[:5])+[None]*max(0,5-len(row))
            if any(isinstance(v,str) and v.startswith('=') for v in values):
                raise ValidationError('Fórmulas não são aceitas nesta demonstração.')
            if tuple(values)==COLUMNS:header_seen=True;continue
            if not header_seen:
                if values[0] and str(values[0]).startswith('Grupo:'):continue
                if not any(v is not None for v in values):continue
                raise ValidationError('Cabeçalhos inesperados.')
            if not any(v is not None and str(v).strip() for v in values):continue
            if str(values[0] or '').startswith(('Grupo:','Resumo do grupo')):continue
            records.append({key:str(value).strip() if value is not None else '' for key,value in zip(COLUMNS,values)})
    finally:book.close()
    if not header_seen:raise ValidationError('Cabeçalhos obrigatórios ausentes.')
    seen=set()
    for record in records:
        ident=record['ID']
        if ident and ident in seen:raise ValidationError('ID preenchido duplicado; destino preservado.')
        if ident:seen.add(ident)
        if not record['Material']:raise ValidationError('Material obrigatório ausente.')
    return records

def publish_local(records:list[dict],target:Path,expected_hash:str|None):
    """Backup conferido e proteção contra mudança do destino após a conferência."""
    current=digest(target) if target.exists() else None
    if current!=expected_hash:raise ValidationError('Destino mudou depois da conferência; escrita bloqueada.')
    target.parent.mkdir(parents=True,exist_ok=True)
    backup=None
    if current:
        folder=target.parent/'backups';folder.mkdir(exist_ok=True)
        backup=folder/f"backup_{datetime.now(timezone.utc).strftime('%Y%m%dT%H%M%S%fZ')}.json"
        backup.write_bytes(target.read_bytes())
        if digest(backup)!=current:raise ValidationError('Backup não confere; escrita bloqueada.')
    # A proteção é local e não representa uma transação remota do Google Sheets.
    body=json.dumps({'demonstracao':True,'colunas':list(COLUMNS),'registros':records},ensure_ascii=False,indent=2)+'\n'
    with tempfile.NamedTemporaryFile('w',encoding='utf-8',dir=target.parent,delete=False,suffix='.tmp') as stream:
        stream.write(body);candidate=Path(stream.name)
    try:
        if (digest(target) if target.exists() else None)!=expected_hash:
            raise ValidationError('Destino mudou antes da escrita; candidato descartado.')
        candidate.replace(target)
        if json.loads(target.read_text(encoding='utf-8'))['registros']!=records:
            raise ValidationError('Releitura não confere; verificar backup antes de nova tentativa.')
    finally:
        if candidate.exists():candidate.unlink()
    return backup

def run(folder:Path,scenario='valido',apply=False):
    folder=folder.resolve();folder.mkdir(parents=True,exist_ok=True)
    source=folder/f'exportacao_{scenario}.xlsx';target=folder/'google_sheets_simulado.json'
    make_sample(source,duplicate=scenario=='duplicado')
    before=digest(target) if target.exists() else None
    records=read_export(source)
    backup=publish_local(records,target,before) if apply else None
    return {'demonstracao':True,'status':'aplicado_localmente' if apply else 'conferencia_aprovada','registros':len(records),'sem_id':sum(not r['ID'] for r in records),'backup':str(backup) if backup else None,'destino_alterado':apply}

if __name__=='__main__':
    parser=argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--cenario',choices=['valido','duplicado'],default='valido')
    parser.add_argument('--aplicar',action='store_true',help='Grava somente o JSON local desta demonstração.')
    parser.add_argument('--pasta-demo',type=Path,default=Path(__file__).resolve().parent/'demo_output')
    args=parser.parse_args()
    try:print(json.dumps(run(args.pasta_demo,args.cenario,args.aplicar),ensure_ascii=False,indent=2))
    except ValidationError as error:
        print(json.dumps({'demonstracao':True,'status':'bloqueado','motivo':str(error)},ensure_ascii=False));raise SystemExit(2)
